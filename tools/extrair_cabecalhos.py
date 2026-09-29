#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
============================================================================
extrair_cabecalhos.py — Tira uma "impressão digital" da planilha em uso
============================================================================
Gera data/cabecalhos_planilha.json com os rótulos reais de cada aba, para que
tools/verificar_schema.js possa provar que o código continua fiel ao arquivo.

Uso:
    python tools/extrair_cabecalhos.py "C:/caminho/Escala_Total_Anestesia.xlsx"

Requer: pip install openpyxl
============================================================================
"""
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("Instale a dependencia: pip install openpyxl")

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA = os.path.join(RAIZ, "data", "cabecalhos_planilha.json")

# Abas cujo cabeçalho não está na linha 1, ou que têm blocos extras a capturar.
LINHAS_EXTRA = {
    "HORAS": [4],
    "FINANCEIRO": [20],
    "INDICADORES": [23],
}


def linha_rotulos(ws, linha):
    """Devolve os rótulos de uma linha, cortando o rabo de células vazias."""
    vals = []
    for c in range(1, ws.max_column + 1):
        v = ws.cell(linha, c).value
        vals.append(None if v is None else str(v))
    while vals and vals[-1] is None:
        vals.pop()
    return vals


def coluna_a(ws, ate=60):
    """Mapa {numero_da_linha: texto} da coluna A — para rótulos de indicadores."""
    out = {}
    for r in range(1, min(ws.max_row, ate) + 1):
        v = ws.cell(r, 1).value
        if v is not None and not str(v).startswith("="):
            out[str(r)] = str(v)
    return out


def main():
    if len(sys.argv) < 2:
        sys.exit("Informe o caminho da planilha .xlsx")
    caminho = sys.argv[1]
    if not os.path.exists(caminho):
        sys.exit("Planilha nao encontrada: " + caminho)

    wb = openpyxl.load_workbook(caminho, data_only=True)
    saida = {}
    for nome in wb.sheetnames:
        ws = wb[nome]
        info = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "cabecalho": linha_rotulos(ws, 1),
            "colunaA": coluna_a(ws),
        }
        for linha in LINHAS_EXTRA.get(nome, []):
            info["linha%d" % linha] = linha_rotulos(ws, linha)
        saida[nome] = info

    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    with open(SAIDA, "w", encoding="utf-8") as f:
        json.dump(saida, f, ensure_ascii=False, indent=1)
    print("Gravado: %s" % SAIDA)
    print("Abas capturadas: %d" % len(saida))


if __name__ == "__main__":
    main()
